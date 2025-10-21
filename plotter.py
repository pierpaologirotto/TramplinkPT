from pathlib import Path
import argparse
import logging
from typing import List, Optional
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator
import openpyxl
import sys

# Aggiunge la root del progetto al percorso
sys.path.append(str(Path(__file__).resolve().parents[2]))
from modules.grafici.table import create_summary_table
from config import paths

# Configura logging leggibile
logging.basicConfig(level=logging.INFO, format="%(message)s")

# Palette di colori da ciclare
COLORS = ['red', 'blue', 'green', 'orange', 'purple', 'cyan', 'magenta', 'brown']


# ===============================================================
# 🔹 FUNZIONI DI SUPPORTO
# ===============================================================
def find_ar_files(base_dir: Path) -> List[Path]:
    """Cerca tutti i file AR-UT*.xlsx sotto la cartella base."""
    files = [p for p in Path(base_dir).rglob("AR-UT*.xlsx") if p.is_file()]
    if not files:
        logging.warning(f"Nessun file AR-UT trovato in {base_dir}")
    else:
        logging.info(f"Trovati {len(files)} file AR-UT in {base_dir}")
    return files


def col_to_float(cell_value):
    """Converte una cella in float, se possibile."""
    if cell_value is None:
        return None
    s = str(cell_value).replace(',', '.').strip()
    try:
        return float(s)
    except ValueError:
        return None


def read_values_from_workbook(path: Path):
    """Legge i valori delle colonne B (etichette) e J (valori numerici) da tutti i fogli."""
    xs: List[int] = []
    ys: List[float] = []
    labels: List[str] = []
    row_indices: List[int] = []

    wb = openpyxl.load_workbook(path, data_only=True)

    for x_index, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        max_row = ws.max_row

        for r in range(2, max_row + 1):
            b_val = ws[f"B{r}"].value
            j_val = col_to_float(ws[f"J{r}"].value)
            if j_val is not None:
                xs.append(x_index)
                ys.append(j_val)
                labels.append(str(b_val) if b_val is not None else "")
                row_indices.append(r)

    wb.close()
    return xs, ys, labels, row_indices


def get_summary_values(path: Path):
    """Estrae gli ultimi valori non vuoti da J, L, O."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    max_row = ws.max_row

    def last_nonempty(col_letter):
        for r in range(max_row, 1, -1):
            val = ws[f"{col_letter}{r}"].value
            if val not in (None, ""):
                return str(val).replace(",", ".")
        return "—"

    values = {
        "Altezza": last_nonempty("J"),
        "Spessore": last_nonempty("L"),
        "qr": last_nonempty("O"),
    }
    wb.close()
    return values


# ===============================================================
# 🔹 PLOT SINGOLO FILE
# ===============================================================
def plot_scatter_for_file(path: Path, show: bool = False) -> Optional[Path]:
    """Genera un grafico scatter per un file AR-UT e lo salva come PNG."""
    xs, ys, labels, row_indices = read_values_from_workbook(path)
    logging.info(f"[{path.name}] → trovati {len(xs)} punti validi per il grafico.")

    if not xs:
        logging.warning(f"Nessun dato valido trovato in {path.name} — nessun grafico generato.")
        return None

    # Crea figura con due assi affiancati (grafico + tabella)
    fig = plt.figure(figsize=(15, 8))
    ax_plot = fig.add_axes([0.05, 0.1, 0.75, 0.8])  # grafico
    ax_table = fig.add_axes([0.82, 0.1, 0.15, 0.8])  # tabella

    # GRAFICO
    point_colors = [COLORS[(r - 2) % len(COLORS)] for r in row_indices]
    ax_plot.scatter(xs, ys, s=40, alpha=0.8, c=point_colors)

    for x_val, y_val, lab in zip(xs, ys, labels):
        if lab:
            ax_plot.annotate(lab, (x_val, y_val), xytext=(3, 3),
                             textcoords='offset points', fontsize=8)

    ax_plot.set_xticks(list(range(1, len(set(xs)) + 1)))
    ax_plot.xaxis.set_major_locator(MaxNLocator(integer=True))
    ax_plot.set_xlabel('Numero foglio')
    ax_plot.set_ylabel('Valori colonna J')
    ax_plot.set_title(f"Scatter AR - {path.stem}")
    ax_plot.grid(axis='y', linestyle='--', alpha=0.4)

    # TABELLA (usa il modulo table.py)
    create_summary_table(ax_table, excel_path=path)

    # Salvataggio
    out_path = path.parent / f"{path.stem}_scatter.png"
    fig.savefig(out_path, dpi=150, bbox_inches='tight')
    plt.close(fig)

    logging.info(f"✅ Grafico salvato in: {out_path}")
    if show:
        plt.show()

    return out_path


# ===============================================================
# 🔹 PLOT DI TUTTI I FILE
# ===============================================================
def plot_all(base_dir: Path, show: bool = False) -> List[Path]:
    """Genera grafici per tutti i file AR-UT nella cartella base."""
    files = find_ar_files(base_dir)
    saved = []

    for f in files:
        try:
            img = plot_scatter_for_file(f, show=show)
            if img:
                saved.append(img)
        except Exception as e:
            logging.exception(f"❌ Errore generando grafico per {f.name}: {e}")

    if saved:
        logging.info(f"\n✅ Totale grafici generati: {len(saved)}")
        for s in saved:
            logging.info(f"   → {s}")
    else:
        logging.warning("⚠️ Nessun grafico generato (nessun dato valido trovato).")

    return saved


# ===============================================================
# 🔹 MAIN
# ===============================================================
if __name__ == '__main__':
    # Usa la stessa cartella base definita in config.paths
    from config import paths

    base = Path(paths.DEST_BASE_1)
    print(f"\n🔍 Cerco file AR-UT in: {base}\n")

    generated = plot_all(base, show=True)

    if generated:
        print(f"\n✅ Generati {len(generated)} grafici:")
        for img in generated:
            print(f"   → {img}")
    else:
        print("\n⚠️ Nessun grafico generato.")
