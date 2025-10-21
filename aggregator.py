# modules/aggregator.py — versione compatibile con .xls
import os
import pandas as pd
from openpyxl import Workbook, load_workbook
from config import paths


HEADER = [
    "NN", "Measurement date", "Measurement time", "Wheelset", "Car No", "Series", "Axle No", "Operator", "Mileage",
    "Height (sH)  (Left)", "Height (sH)  (Right)", "Thickness (sD)  (Left)", "Thickness (sD) (Right)",
    "Parameter (sF)  (Left)", "Parameter (sF) (Right)", "Thickness (sDF)  (Left)", "Thickness (sDF) (Right)",
    "Angle (A)  (Left)", "Angle (A) (Right)"
]
MAX_COLS = len(HEADER)


def ensure_ar_file(ut_folder):
    ar_filename = f"AR-{os.path.basename(ut_folder)}.xlsx"
    ar_path = os.path.join(ut_folder, ar_filename)
    if os.path.exists(ar_path):
        return ar_path
    wb = Workbook()
    ws = wb.active
    ws.title = "1"
    ws.append(HEADER)
    for i in range(2, 7):
        ws_i = wb.create_sheet(str(i))
        ws_i.append(HEADER)
    wb.save(ar_path)
    wb.close()
    print(f"[CREATO] {ar_path}")
    return ar_path


def convert_xls_to_xlsx(src_path):
    """Converte un file .xls in .xlsx (temporaneo) e restituisce il nuovo path"""
    tmp_path = src_path + "_tmp.xlsx"
    try:
        df = pd.read_excel(src_path, header=None, engine="xlrd")
        df.to_excel(tmp_path, header=False, index=False, engine="openpyxl")
        print(f"  [CONVERTITO] {os.path.basename(src_path)} -> {os.path.basename(tmp_path)}")
        return tmp_path
    except Exception as e:
        print(f"  [ERRORE conversione {src_path}: {e}]")
        return None


def read_row_values(src_path, row_num):
    """Legge una riga (row_num) dal primo foglio del file .xlsx"""
    try:
        wb = load_workbook(src_path, data_only=True)
        ws = wb.worksheets[0]
        if ws.max_row < row_num:
            wb.close()
            return ["nessun dato"] * MAX_COLS
        vals = []
        for c in range(1, MAX_COLS + 1):
            v = ws.cell(row=row_num, column=c).value
            vals.append(v if v is not None else "nessun dato")
        wb.close()
        return vals
    except Exception as e:
        print(f"  [ERRORE lettura {src_path}: {e}]")
        return ["nessun dato"] * MAX_COLS


def collect_data_for_ut(ut_folder):
    raccolta = os.path.join(ut_folder, "Raccolta")
    if not os.path.isdir(raccolta):
        print(f"[SKIP] {ut_folder}: no 'Raccolta'")
        return

    print(f"[START UT] {ut_folder}")
    ar_path = ensure_ar_file(ut_folder)
    wb_target = load_workbook(ar_path)

    files = sorted([f for f in os.listdir(raccolta)
                    if (f.lower().endswith(".xlsx") or f.lower().endswith(".xls")) and not f.startswith("~$")])
    if not files:
        print(f"[INFO] {ut_folder}: Raccolta vuota")
        wb_target.close()
        return

    for filename in files:
        src_path = os.path.join(raccolta, filename)
        if filename.lower().endswith(".xls"):
            src_path_conv = convert_xls_to_xlsx(src_path)
            if not src_path_conv:
                continue
            src_path = src_path_conv  # sostituisce il path con quello convertito

        print(f"  -> Leggo {filename}")
        for sheet_idx, src_row in enumerate(range(2, 8), start=1):
            vals = read_row_values(src_path, src_row)
            ws_t = wb_target[str(sheet_idx)]
            ws_t.append(vals)

        # elimina file temporaneo se era .xls
        if filename.lower().endswith(".xls") and os.path.exists(src_path):
            os.remove(src_path)

    wb_target.save(ar_path)
    wb_target.close()
    print(f"[SAVED] {ar_path}\n")


def aggregate_all_ut():
    base = paths.DEST_BASE_1
    if not os.path.isdir(base):
        print(f"[ERROR] DEST_BASE_1 not found: {base}")
        return
    ut_folders = [d for d in sorted(os.listdir(base))
                  if os.path.isdir(os.path.join(base, d)) and d.startswith("UT ")]
    if not ut_folders:
        print("[INFO] No UT folders found.")
        return
    for ut in ut_folders:
        collect_data_for_ut(os.path.join(base, ut))
    print("✅ Aggregazione completata.")
